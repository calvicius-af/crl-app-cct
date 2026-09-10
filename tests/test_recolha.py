"""Testes da recolha do BTE (SPEC-0001, ADR-0014).

Nenhum destes testes toca na rede: o transporte é injetado em `abridor=`.
O único que a usa está protegido por CCT_TESTE_REDE=1.

O índice de ensaio é o do **BTE n.º 31 de 2026** — cabeçalho e linhas reais,
reconstruídos com openpyxl para não versionar binários no repositório.
"""
import os
from pathlib import Path

import pytest

from cct import recolha
from cct.recolha import (ErroRede, Registo, _nome_seguro, descarregar_item,
                         familia, ler_indice, recolher, validar_url)

CABECALHO = [
    "ANO", "ID:", "TITULO DO DOCUMENTO:", "TIPO DE DOCUMENTO:",
    "TIPO DE DOCUMENTO:", "Nº VOLUME DO BOLETIM:", "Nº DO BOLETIM:",
    "DATA DO BOLETIM:", "DATA DE DISTRIBUIÇÃO DO BOLETIM:",
    "PÁGINA NA VERSÃO ESCRITA:", "CAE:", "COD:\n(IRCT)",
    "VIDE DOCUMENTO(S) EM VIGOR:", "LINK VIDE DOCUMENTO(S) EM VIGOR:",
    "DOC(S) ALTERADO(S) POR ESTE:", "DOC(S) ALTERADO(S) POR ESTE:2",
    "DOC(S) QUE ALTERA(M) ESTE:", "SECTOR(ES) DE ACTIVIDADE:", "OUTORGANTE(S):",
    "fich+X310+X1:AB35", "Página (criado)", "Link para o documento (CRIADO)",
]

URL = "https://bte.dgcp.mtsss.gov.pt/documentos/2026/31"

LINHAS = [
    ("377/2026", "Contrato coletivo entre a Associação do Comércio e Serviços da "
     "Região do Algarve - ACRAL e o CESP - Sindicato dos Trabalhadores do Comércio",
     "CCT", "27251",
     "Associação do Comércio e Serviços da Região do Algarve - ACRAL; "
     "CESP - Sindicato dos Trabalhadores do Comércio, Escritórios e Serviços de Portugal",
     "00260057.pdf"),
    ("378/2026", "Contrato coletivo entre a Confederação Nacional das Instituições "
     "de Solidariedade - CNIS e a Federação Nacional dos Sindicatos",
     "CCT-ALT", "26760",
     "Confederação Nacional das Instituições de Solidariedade - CNIS; "
     "Federação Nacional dos Sindicatos dos Trabalhadores em Funções Públicas "
     "e Sociais - FNSTFPS",
     "00580059.pdf"),
    ("379/2026", "Contrato coletivo entre a Associação das Empresas de Vinho do "
     "Porto (AEVP) e a FESAHT - Federação dos Sindicatos da Agricultura",
     "CCT-ALT", "26651",
     "Associação das Empresas de Vinho do Porto (AEVP); "
     "FESAHT - Federação dos Sindicatos da Agricultura, Alimentação, Bebidas, "
     "Hotelaria e Turismo de Portugal",
     "00600062.pdf"),
    ("382/2026", "Acordo de empresa entre a Empresa Metropolitana de Estacionamento "
     "da Maia, EM e o Sindicato dos Trabalhadores da Administração Pública",
     "AE", "47252",
     "Empresa Metropolitana de Estacionamento da Maia, EM; "
     "Sindicato dos Trabalhadores da Administração Pública e de Entidades com "
     "Fins Públicos - SINTAP",
     "00880122.pdf"),
    ("385/2026", "Acordo de empresa entre a AP Solutions GMBH - Sucursal em Portugal "
     "e o Sindicato dos Trabalhadores da Actividade Seguradora (STAS)",
     "AE-ALT", "47140",
     "AP Solutions GMBH - Sucursal em Portugal; "
     "Sindicato dos Trabalhadores da Actividade Seguradora (STAS)",
     "01500163.pdf"),
    # retificação: a coluna de outorgantes vem vazia no índice real
    ("387/2026", "Acordo de empresa entre a CARRISTUR - Inovação em Transportes "
     "Urbanos e Regionais, Sociedade Unipessoal L.da e a Associação Sindical das "
     "Trabalhadoras e Trabalhadores dos Transportes",
     "AE-ALT-RECT", "47109", None, "01680170.pdf"),
    # linha fora do âmbito, para verificar que não é descarregada em silêncio
    ("391/2026", "Estatutos do Sindicato dos Trabalhadores de Alguma Coisa",
     "ST", "50001", None, "01900191.pdf"),
]

PDF_FALSO = b"%PDF-1.4\n% ficheiro de ensaio\n"


def escrever_indice(pasta: Path, nome: str = "BTE31_2026.xlsx") -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "31"
    ws.append(CABECALHO)
    for id_dgert, titulo, tipo, cod, outorgantes, ficheiro in LINHAS:
        linha = [None] * len(CABECALHO)
        linha[0], linha[1], linha[2], linha[3] = 2026, id_dgert, titulo, tipo
        linha[5], linha[6] = 93, 31
        linha[7], linha[8] = "2026-08-22", "2026-08-24"
        linha[11] = cod
        linha[18] = outorgantes
        linha[20] = ficheiro
        linha[21] = f"{URL}/{ficheiro}"
        ws.append(linha)
    caminho = pasta / nome
    wb.save(caminho)
    return caminho


class AbridorFalso:
    """Transporte de ensaio: conta pedidos e responde como o servidor do BTE."""

    def __init__(self, corpo=PDF_FALSO, estado=200, etag='"abc"'):
        self.corpo, self.estado, self.etag = corpo, estado, etag
        self.pedidos: list[tuple[str, dict]] = []

    def __call__(self, url, cabecalhos=None):
        recolha.validar_url(url)
        self.pedidos.append((url, dict(cabecalhos or {})))
        if cabecalhos and cabecalhos.get("If-None-Match") == self.etag:
            return recolha.Resposta(304, {"ETag": self.etag}, b"")
        return recolha.Resposta(self.estado,
                                {"ETag": self.etag,
                                 "Last-Modified": "Tue, 18 Aug 2026 10:31:46 GMT"},
                                self.corpo)


def recusa_rede(url, cabecalhos=None):
    raise AssertionError(f"pedido de rede não autorizado: {url}")


@pytest.fixture()
def ambiente(tmp_path):
    indice = escrever_indice(tmp_path)
    registo = Registo(tmp_path / "registo.jsonl")
    return indice, registo, tmp_path / "interim"


# ------------------------------------------------------------ leitura do índice

def test_le_o_indice_real_do_bte31(ambiente):
    indice, _, _ = ambiente
    itens = ler_indice(indice)
    assert len(itens) == len(LINHAS)
    primeiro = itens[0]
    assert primeiro["tipo"] == "CCT"
    assert primeiro["ano"] == 2026 and primeiro["num_bte"] == 31
    assert primeiro["cod_irct"] == "27251"
    assert primeiro["url"] == f"{URL}/00260057.pdf"
    assert primeiro["ficheiro"] == "00260057.pdf"
    assert primeiro["chave"] == "2026/31/00260057"
    assert primeiro["familia"] == "convencao"
    assert primeiro["posicao"] == 1


def test_cabecalho_repetido_nao_apaga_o_tipo(ambiente):
    # o índice tem 'TIPO DE DOCUMENTO:' duas vezes, a segunda vazia
    indice, _, _ = ambiente
    assert [i["tipo"] for i in ler_indice(indice)][:3] == ["CCT", "CCT-ALT", "CCT-ALT"]


def test_retificacao_sem_outorgantes_continua_a_ser_lida(ambiente):
    indice, _, _ = ambiente
    rect = [i for i in ler_indice(indice) if i["tipo"] == "AE-ALT-RECT"][0]
    assert rect["outorgantes"] == ""
    assert "CARRISTUR" in rect["titulo"]
    assert rect["familia"] == "convencao"


@pytest.mark.parametrize("tipo,esperado", [
    ("CCT", "convencao"), ("CCT-ALT", "convencao"), ("AE-ALT-RECT", "convencao"),
    ("ACT", "convencao"), ("ACTV", "convencao"),
    ("PE", "extensao"), ("PCT", "extensao"),
    ("AVISO", "aviso"), ("AA", "adesao"),
    ("ST", None), ("", None), ("Eleições", None),
])
def test_familia_por_tipo(tipo, esperado):
    assert familia(tipo) == esperado


# --------------------------------------------------------------- rede e limites

@pytest.mark.parametrize("url", [
    "http://bte.dgcp.mtsss.gov.pt/x.pdf",                  # sem TLS
    "https://exemplo.pt/documentos/2026/31/00260057.pdf",  # anfitrião não permitido
    "https://bte.dgcp.mtsss.gov.pt.exemplo.pt/x.pdf",      # sufixo enganador
    "ficheiro.pdf",
])
def test_url_recusado(url):
    with pytest.raises(ErroRede):
        validar_url(url)


def test_url_permitido():
    assert validar_url(f"{URL}/00260057.pdf")


def test_redireccionamento_para_fora_e_recusado():
    handler = recolha._RedireccionamentoVerificado()
    with pytest.raises(ErroRede):
        handler.redirect_request(None, None, 302, "Found", {},
                                 "https://exemplo.pt/x.pdf")


def test_sem_confirmar_rede_nao_ha_pedidos(ambiente):
    indice, registo, interim = ambiente
    resumo = recolher([indice], interim, registo, rede=False, abridor=recusa_rede)
    assert resumo["pedidos_de_rede"] == 0
    assert resumo["por_estado"]["por_descarregar"] == 6
    assert not list(interim.rglob("*.pdf"))


def test_tipo_fora_do_ambito_e_reportado_nao_descarregado(ambiente):
    indice, registo, interim = ambiente
    abridor = AbridorFalso()
    resumo = recolher([indice], interim, registo, rede=True, abridor=abridor, pausa=0)
    assert resumo["tipos_desconhecidos"] == {"ST": 1}
    assert len(abridor.pedidos) == 6


def test_familias_limitam_o_que_e_descarregado(ambiente):
    indice, registo, interim = ambiente
    abridor = AbridorFalso()
    recolher([indice], interim, registo, rede=True, abridor=abridor, pausa=0,
             familias=("extensao",))
    assert abridor.pedidos == []


# ----------------------------------------------------------------- idempotência

def test_descarrega_uma_vez_e_nao_repete(ambiente):
    indice, registo, interim = ambiente
    abridor = AbridorFalso()

    r1 = recolher([indice], interim, registo, rede=True, abridor=abridor, pausa=0)
    assert r1["por_estado"]["descarregado"] == 6
    assert (interim / "2026" / "31" / "00260057.pdf").read_bytes() == PDF_FALSO

    registo2 = Registo.carregar(registo.caminho)
    r2 = recolher([indice], interim, registo2, rede=True, abridor=recusa_rede, pausa=0)
    assert r2["por_estado"]["ja_existente"] == 6
    assert r2["pedidos_de_rede"] == 0


def test_ficheiro_alterado_no_disco_faz_pedido_condicional(ambiente):
    indice, registo, interim = ambiente
    abridor = AbridorFalso()
    recolher([indice], interim, registo, rede=True, abridor=abridor, pausa=0)

    alvo = interim / "2026" / "31" / "00260057.pdf"
    alvo.write_bytes(b"%PDF-1.4\nconteudo diferente\n")

    abridor2 = AbridorFalso()
    resumo = recolher([indice], interim, Registo.carregar(registo.caminho),
                      rede=True, abridor=abridor2, pausa=0)
    assert resumo["por_estado"]["inalterado"] == 1        # 304 do servidor
    assert abridor2.pedidos[0][1]["If-None-Match"] == '"abc"'


def test_ficheiro_apagado_e_descarregado_de_novo_por_inteiro(ambiente):
    indice, registo, interim = ambiente
    recolher([indice], interim, registo, rede=True, abridor=AbridorFalso(), pausa=0)
    (interim / "2026" / "31" / "00260057.pdf").unlink()

    abridor = AbridorFalso()
    resumo = recolher([indice], interim, Registo.carregar(registo.caminho),
                      rede=True, abridor=abridor, pausa=0)
    assert resumo["por_estado"]["descarregado"] == 1
    assert "If-None-Match" not in abridor.pedidos[0][1]   # pede a cópia inteira


def test_resposta_que_nao_e_pdf_falha_sem_escrever(ambiente):
    indice, registo, interim = ambiente
    resumo = recolher([indice], interim, registo, rede=True, pausa=0,
                      abridor=AbridorFalso(corpo=b"<html>manutencao</html>"))
    assert resumo["por_estado"]["falhado"] == 6
    assert not list(interim.rglob("*.pdf"))
    assert not list(interim.rglob("*.part"))


def test_erro_404_nao_e_repetido_nem_escreve(ambiente):
    indice, registo, interim = ambiente
    abridor = AbridorFalso(estado=404, corpo=b"")
    resumo = recolher([indice], interim, registo, rede=True, abridor=abridor, pausa=0)
    assert resumo["por_estado"]["falhado"] == 6
    assert len(abridor.pedidos) == 6                     # uma tentativa por documento
    assert "HTTP 404" in resumo["problemas"][0]


# ----------------------------------------- persistência periódica do registo

def test_registo_persiste_periodicamente_nao_so_no_fim(ambiente, monkeypatch):
    """Ver PR #35, achado nº6: o registo não pode ficar só à espera do fim
    da corrida para ser gravado — senão uma interrupção a meio perde tudo."""
    indice, registo, interim = ambiente
    monkeypatch.setattr(recolha, "INTERVALO_PERSISTENCIA", 2)

    gravacoes = []
    original = registo.guardar

    def guardar_e_contar():
        original()
        gravacoes.append(len(Registo.carregar(registo.caminho).entradas))

    registo.guardar = guardar_e_contar
    recolher([indice], interim, registo, rede=True, abridor=AbridorFalso(), pausa=0)

    # 6 documentos válidos, intervalo=2 → pelo menos 3 gravações intermédias
    # (a última coincide com a gravação final do `finally`)
    assert len(gravacoes) >= 3


def test_interrupcao_a_meio_nao_perde_o_que_ja_foi_descarregado(ambiente, monkeypatch):
    monkeypatch.setattr(recolha, "INTERVALO_PERSISTENCIA", 2)
    indice, registo, interim = ambiente

    class AbridorComFalha(AbridorFalso):
        def __init__(self):
            super().__init__()
            self.chamadas = 0

        def __call__(self, url, cabecalhos=None):
            self.chamadas += 1
            if self.chamadas > 3:
                raise RuntimeError("falha simulada a meio da corrida")
            return super().__call__(url, cabecalhos)

    with pytest.raises(RuntimeError):
        recolher([indice], interim, registo, rede=True,
                abridor=AbridorComFalha(), pausa=0)

    # apesar da excepção não tratada, o `finally` já gravou o que foi feito
    relido = Registo.carregar(registo.caminho)
    descarregados = [e for e in relido.entradas.values()
                     if (e.get("descarga") or {}).get("estado") == "descarregado"]
    assert len(descarregados) >= 2

    # e uma corrida seguinte não repete os pedidos já bem sucedidos
    resumo2 = recolher([indice], interim, relido, rede=True,
                       abridor=AbridorFalso(), pausa=0)
    assert resumo2["pedidos_de_rede"] < 6


def test_nao_ficam_ficheiros_temporarios(ambiente):
    indice, registo, interim = ambiente
    recolher([indice], interim, registo, rede=True, abridor=AbridorFalso(), pausa=0)
    assert not list(interim.rglob("*.part"))


# ------------------------------------------------------- travessia de diretório

@pytest.mark.parametrize("bruto,esperado", [
    ("../../../../tmp/evil.pdf", "evil.pdf"),        # neutralizado: só o nome sobra
    ("../../../etc/passwd", "passwd"),
    ("/etc/passwd", "passwd"),
    ("00260057.pdf", "00260057.pdf"),                # caso normal, sem alteração
])
def test_nome_seguro_neutraliza_travessia_para_o_nome_simples(bruto, esperado):
    assert _nome_seguro(bruto) == esperado


@pytest.mark.parametrize("bruto", ["..", ".", ""])
def test_nome_seguro_recusa_o_que_nao_sobra_nome_nenhum(bruto):
    if bruto == "":
        assert _nome_seguro(bruto) == ""    # ausência de valor é válida (RF opcional)
    else:
        with pytest.raises(ValueError):
            _nome_seguro(bruto)


def test_indice_com_ficheiro_de_travessia_e_neutralizado_nao_escreve_fora_do_destino(
        ambiente):
    indice, registo, interim = ambiente
    # injecta um valor de travessia diretamente na folha, como viria de uma
    # célula corrompida ou adulterada do índice
    import openpyxl
    wb = openpyxl.load_workbook(indice)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[20].value == "00260057.pdf":            # coluna "Página (criado)"
            row[20].value = "../../../../../../tmp/evil_travessia.pdf"
            row[21].value = f"{URL}/00260057.pdf"       # URL continua válido
            break
    wb.save(indice)

    resumo = recolher([indice], interim, registo, rede=True,
                      abridor=AbridorFalso(), pausa=0)

    # os 6 documentos continuam a ser processados — a travessia foi apenas
    # reduzida ao nome simples "evil_travessia.pdf", dentro do destino certo
    assert resumo["pedidos_de_rede"] == 6
    assert resumo["por_estado"].get("descarregado", 0) == 6
    alvo = interim / "2026" / "31" / "evil_travessia.pdf"
    assert alvo.exists()
    # nada foi escrito fora de data/interim/recolha
    for caminho in interim.rglob("*"):
        assert interim.resolve() in caminho.resolve().parents or caminho == interim
    # em particular, nada foi escrito para fora do sistema de ficheiros de ensaio
    assert not (interim.parent.parent / "evil_travessia.pdf").exists()


def test_descarregar_item_recusa_caminho_fora_do_destino_mesmo_sem_completar(
        ambiente):
    """Cinto-e-suspensórios: mesmo chamando a função directamente, sem passar
    por _completar()/_nome_seguro(), o caminho de escrita nunca escapa do
    destino."""
    _indice, registo, interim = ambiente
    item = {"chave": "x", "ano": 2026, "num_bte": 31,
           "ficheiro": "../../../../evil.pdf",
           "url": f"{URL}/00260057.pdf"}
    with pytest.raises(ValueError, match="fora de"):
        descarregar_item(item, interim, registo, abridor=AbridorFalso())


# --------------------------------------------------- validação do registo

def test_registo_carrega_ignorando_linha_invalida(ambiente, capsys):
    """Ver PR #35, achado nº10: uma linha malformada não pode contaminar
    o resto do registo nem propagar valores errados para cct.nomeacao."""
    indice, registo, interim = ambiente
    recolher([indice], interim, registo, rede=True, abridor=AbridorFalso(), pausa=0)

    with open(registo.caminho, "a", encoding="utf-8") as f:
        f.write('{"chave": "invalida/sem/ano", "ano": "não é um número"}\n')
        f.write("isto nem sequer é JSON\n")
        f.write('{"num_bte": 31}\n')             # sem "chave" — required

    capsys.readouterr()  # limpa o que já foi impresso
    relido = Registo.carregar(registo.caminho)
    saida = capsys.readouterr().out
    assert saida.count("AVISO —") == 3
    assert "invalida/sem/ano" not in relido.entradas
    # as entradas válidas (incluindo a "ignorada" por família fora de âmbito)
    # continuam todas presentes — só as 3 linhas injectadas foram rejeitadas
    assert len(relido.entradas) == len(LINHAS)


def test_registo_sobrevive_a_uma_ida_ao_disco(ambiente):
    indice, registo, interim = ambiente
    recolher([indice], interim, registo, rede=True, abridor=AbridorFalso(), pausa=0)
    relido = Registo.carregar(registo.caminho)
    entrada = relido.get("2026/31/00260057")
    assert entrada["cod_irct"] == "27251"
    assert entrada["descarga"]["sha256"]
    assert entrada["descarga"]["bytes"] == len(PDF_FALSO)


@pytest.mark.skipif(os.environ.get("CCT_TESTE_REDE") != "1",
                    reason="teste com rede (correr com CCT_TESTE_REDE=1)")
def test_descarga_real_do_bte(tmp_path):
    indice = escrever_indice(tmp_path)
    registo = Registo(tmp_path / "registo.jsonl")
    resumo = recolher([indice], tmp_path / "interim", registo, rede=True,
                      pausa=0.5, limite=1)
    assert resumo["por_estado"].get("descarregado") == 1
    pdf = next((tmp_path / "interim").rglob("*.pdf"))
    assert pdf.read_bytes().startswith(b"%PDF")
